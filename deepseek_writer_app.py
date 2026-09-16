import re
import streamlit as st
from docx import Document
from io import BytesIO
from pathlib import Path
from openai import OpenAI


# ============== 工具函数 ==============

def count_chinese_words(text):
    """统计中文字符数（不含标点和空白）"""
    return len(re.findall(r"[\u4e00-\u9fff]", text))


def load_docx_text(file_bytes):
    """读取 docx 文件为段落列表"""
    doc = Document(BytesIO(file_bytes))
    return [p.text.strip() for p in doc.paragraphs]


def list_repo_docs():
    """列出仓库 docs/ 目录下的文档"""
    docs_dir = Path(__file__).parent / "docs"
    if not docs_dir.exists():
        return []
    return sorted([f.name for f in docs_dir.iterdir() if f.suffix.lower() in (".docx", ".md", ".txt")])


def load_repo_doc(filename):
    """读取仓库 docs/ 目录下的 docx 文件"""
    file_path = Path(__file__).parent / "docs" / filename
    if file_path.exists():
        return file_path.read_bytes()
    return None


def load_repo_doc_text(filename):
    """读取仓库 docs/ 目录下的 docx/md/txt 文件，返回文本"""
    file_path = Path(__file__).parent / "docs" / filename
    if not file_path.exists():
        return ""
    suffix = file_path.suffix.lower()
    if suffix == ".docx":
        return "\n".join([p.text for p in Document(BytesIO(file_path.read_bytes())).paragraphs])
    elif suffix in (".md", ".txt"):
        return file_path.read_text(encoding="utf-8")
    return ""


def load_repo_knowledge_text(filename):
    """读取仓库根目录下的 md 知识库文件"""
    file_path = Path(__file__).parent / filename
    if file_path.exists():
        return file_path.read_text(encoding="utf-8")
    return ""


def num_to_chinese(n):
    """数字转中文，支持 1-99"""
    digits = "零一二三四五六七八九"
    if n < 10:
        return digits[n]
    elif n < 20:
        return "十" + (digits[n - 10] if n > 10 else "")
    else:
        tens = n // 10
        units = n % 10
        if units == 0:
            return digits[tens] + "十"
        else:
            return digits[tens] + "十" + digits[units]


def extract_chapter(paragraphs, chapter_num):
    """提取指定章节的正文和细纲"""
    chapter_num_cn = num_to_chinese(chapter_num)

    # 找正文起点标记
    body_start = 0
    for i, t in enumerate(paragraphs):
        if "以上为番茄版大纲" in t or t == "正文":
            body_start = i
            break

    # 正文
    body_texts = []
    body_started = False
    chapter_pattern = re.compile(rf"^第({chapter_num}|{chapter_num_cn})章")
    next_pattern = re.compile(r"^第(\d+|[一二三四五六七八九十百]+)章")
    for i in range(body_start, len(paragraphs)):
        t = paragraphs[i]
        if not t:
            continue
        if chapter_pattern.search(t):
            body_started = True
            body_texts.append(t)
            continue
        if body_started:
            if next_pattern.match(t):
                break
            body_texts.append(t)

    # 细纲（默认在前面的段落里找）
    outline_texts = []
    outline_started = False
    outline_pattern = re.compile(rf"^\s*Ch{chapter_num}[^0-9]|^\s*第({chapter_num}|{chapter_num_cn})章")
    next_outline_pattern = re.compile(r"^\s*Ch\d+[^0-9]|^\s*第(\d+|[一二三四五六七八九十百]+)章")
    for i, t in enumerate(paragraphs):
        if outline_pattern.search(t):
            outline_started = True
            outline_texts.append(t)
            continue
        if outline_started:
            if next_outline_pattern.match(t):
                break
            outline_texts.append(t)

    return "\n".join(body_texts), "\n".join(outline_texts)


def extract_outline_from_text(text, chapter_num):
    """从独立细纲文本中提取指定章节细纲"""
    chapter_num_cn = num_to_chinese(chapter_num)
    lines = text.split("\n")
    result = []
    started = False
    pattern = re.compile(rf"^\s*Ch{chapter_num}[^0-9]|^\s*第({chapter_num}|{chapter_num_cn})章")
    next_pattern = re.compile(r"^\s*Ch\d+[^0-9]|^\s*第(\d+|[一二三四五六七八九十百]+)章")
    for line in lines:
        t = line.strip()
        if not t:
            continue
        if pattern.search(t):
            started = True
            result.append(t)
            continue
        if started:
            if next_pattern.match(t):
                break
            result.append(t)
    return "\n".join(result)


def call_deepseek(api_key, system_prompt, user_prompt, model="deepseek-chat", temperature=0.7):
    """调用 DeepSeek API"""
    client = OpenAI(api_key=api_key, base_url="https://api.deepseek.com")
    response = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        temperature=temperature,
        stream=False,
    )
    return response.choices[0].message.content


# ============== Prompt 构建 ==============

def build_system_prompt(chapter6_sample, methodology, ch1_case=""):
    ch1_section = f"\n\n下面是第1章去AI味批改案例，请学习其中的修改原则和自检三问：\n\n{ch1_case}" if ch1_case else ""
    return f"""你是一位专业悬疑/克系小说编辑与写手，熟悉番茄小说风格，擅长去除AI味。
你的任务是根据用户提供的《不周》章节，进行改写或扩写，使其语言自然、节奏紧凑、五感真实、人物对话符合身份，彻底消除AI写作痕迹。

核心要求：
1. 写景必须服务于人物情绪和剧情推进，删掉无关冗余细节。
2. 用动作和微表情代替刻意的"沉默"和"留白"。
3. 对话符合人物关系，真实口语，有潜台词。
4. 伏笔落地，让异常有真实触感（温度、触感、气味、声音）。
5. 每章结尾尽量落在画面或动作上，不用解释性总结。
6. 章末附字数和词频自检（列出本章出现频率最高的3个词，并说明是否重复）。

下面是第6章改写范本，请严格参照其语感、节奏和细节密度：

{chapter6_sample}{ch1_section}

以下是去AI味方法论补充：
{methodology}
"""


def build_user_prompt(chapter_num, mode, body, outline, feedback=""):
    if not body.strip() and not outline.strip():
        raise ValueError(
            f"未找到第{chapter_num}章内容，请检查源文件章节标题是否为「第{chapter_num}章」或「第{num_to_chinese(chapter_num)}章」格式"
        )

    feedback_instruction = ""
    if feedback and feedback.strip():
        feedback_instruction = f"\n\n额外修改意见（请重点按以下意见调整）：\n{feedback}\n"

    if mode == "改写":
        return f"""请改写《不周》第{chapter_num}章。

【当前正文】：
{body if body else '（当前正文为空）'}

【本章细纲】：
{outline if outline else '（未提供细纲）'}

要求：
- 严格按细纲归位剧情
- 字数 2200-2800 字
- 彻底去 AI 味
- 参照第6章范本风格
- 章末附字数和词频自检{feedback_instruction}

请直接输出改写后的完整章节。"""
    else:
        return f"""请扩写《不周》第{chapter_num}章。

【本章细纲】：
{outline if outline else '（未提供细纲）'}

【当前正文】（如有）：
{body if body else '（当前正文为空）'}

要求：
- 严格按细纲扩写，不偏离主线
- 字数 2200-2800 字
- 彻底去 AI 味
- 参照第6章范本风格
- 章末附字数和词频自检{feedback_instruction}

请直接输出扩写后的完整章节。"""


def build_outline_prompt(outline_content, volume, chapter_count):
    return f"""请根据以下《不周》总大纲，生成第{volume}卷的逐章细纲，共{chapter_count}章。

【总大纲】：
{outline_content}

要求：
1. 每章一条，格式：第X章 标题 + 300-500 字细纲
2. 每章必须包含：场景、核心冲突、人物动作、悬念钩子
3. 保持悬疑节奏，每3-5章一个小高潮
4. 去AI味：避免模板化过渡句、堆砌形容词、解释性总结
5. 细纲语言简洁，留给正文写作空间

请直接输出第{volume}卷完整细纲。"""


def build_continuation_prompt(chapter_num, prev_texts, outline, ch1_case=""):
    prev_section = "\n\n".join(prev_texts) if prev_texts else "（无前文）"
    ch1_section = f"\n\n【Ch1 去AI味案例】：\n{ch1_case}\n请严格按其中原则写作。" if ch1_case else ""
    return f"""请为《不周》续写第{chapter_num}章。

【前文节选】：
{prev_section}

【本章细纲】：
{outline if outline else '（未提供细纲，请根据前文合理推进）'}
{ch1_section}

要求：
1. 承接前文人物状态和悬念，保持风格一致
2. 严格按细纲推进，可适度发挥但主线不跑偏
3. 字数 2200-2800 字
4. 彻底去 AI 味：写景服务情绪、删冗余、动作代沉默、对话真实、伏笔有触感
5. 章末附字数和词频自检
6. 只输出本章正文，不要总结全书

请直接输出第{chapter_num}章完整正文。"""


# ============== 知识库加载 ==============
@st.cache_resource
def load_knowledge():
    sample_path = Path(__file__).parent / "不周_Ch6_改写样章.md"
    methodology_path = Path(__file__).parent / "去AI味长篇小说写作方法论_report.md"
    ch1_case_path = Path(__file__).parent / "不周_Ch1_去AI味案例.md"
    chapter6_sample = sample_path.read_text(encoding="utf-8") if sample_path.exists() else "（未找到第6章范本）"
    methodology = methodology_path.read_text(encoding="utf-8") if methodology_path.exists() else ""
    ch1_case = ch1_case_path.read_text(encoding="utf-8") if ch1_case_path.exists() else ""
    return build_system_prompt(chapter6_sample, methodology, ch1_case)


# ============== 页面 UI ==============

st.title("✍️ 墨笔·不周写作助手")
st.caption("基于 DeepSeek API 的《不周》改写 / 扩写 / 续写工具")

# 侧边栏配置
with st.sidebar:
    st.header("⚙️ 配置")
    api_key = st.text_input("DeepSeek API Key", type="password", help="在 platform.deepseek.com 获取")
    model = st.selectbox("模型", ["deepseek-chat", "deepseek-reasoner"], index=0)
    temperature = st.slider("Temperature", 0.0, 1.0, 0.7, 0.05)

    st.divider()
    st.markdown("""
    **使用步骤：**
    1. 填入 API Key
    2. 在「改写/扩写」或「续写」标签选择源文件
    3. 选择章节号和模式
    4. 点击运行
    """)

# 标签页
tab1, tab2, tab3, tab4 = st.tabs(["改写/扩写", "创作大纲", "生成细纲", "续写"])

system_prompt = load_knowledge()

# ============== Tab 1: 改写/扩写 ==============
with tab1:
    st.header("1. 选择源文件")
    source_option = st.radio("文件来源", ["从仓库选择", "本地上传"], index=0, horizontal=True, key="t1_source")

    paragraphs = None
    if source_option == "从仓库选择":
        repo_files = list_repo_docs()
        if repo_files:
            selected_file = st.selectbox("仓库文档列表", repo_files, help="把 docx 文件放进仓库的 docs/ 目录，即可在此选择", key="t1_repo")
            if selected_file:
                file_bytes = load_repo_doc(selected_file)
                if file_bytes:
                    paragraphs = load_docx_text(file_bytes)
                    st.success(f"已加载：{selected_file}，共 {len(paragraphs)} 段")
        else:
            st.warning("仓库里还没有文档。请在仓库根目录新建 `docs/` 文件夹，把 docx/md/txt 放进去。")
    else:
        uploaded_file = st.file_uploader("上传《不周》docx 文件", type=["docx"], key="t1_upload")
        if uploaded_file is not None:
            paragraphs = load_docx_text(uploaded_file.getvalue())
            st.success(f"文件已读取，共 {len(paragraphs)} 段")

    if paragraphs is not None:
        chapter_numbers = sorted(set(
            int(m.group(1))
            for t in paragraphs
            for m in [re.search(r"第(\d+)章", t)]
            if m
        ))
        if chapter_numbers:
            st.info(f"检测到章节：第 {chapter_numbers[0]} 章 至 第 {chapter_numbers[-1]} 章")

        # 单章改写
        st.header("2. 单章改写")
        col1, col2 = st.columns(2)
        with col1:
            single_chapter = st.number_input("章节号", min_value=1, max_value=200, value=7, step=1, key="single_ch")
        with col2:
            single_mode = st.selectbox("模式", ["改写", "扩写"], index=0, key="single_mode")

        single_feedback = st.text_area(
            "修改意见（可选）",
            placeholder="例如：加强陆潮的紧张感，减少环境描写，老周台词要更糙...",
            key="single_feedback",
        )

        if st.button("🚀 开始生成", type="primary", key="single_run"):
            if not api_key:
                st.error("请先填写 DeepSeek API Key")
            else:
                with st.spinner("正在提取章节内容..."):
                    body, outline = extract_chapter(paragraphs, int(single_chapter))

                if not outline:
                    st.warning("未找到该章节细纲，将直接按原文改写")
                if not body:
                    st.warning("未找到该章节正文，将按细纲扩写")

                try:
                    user_prompt = build_user_prompt(int(single_chapter), single_mode, body, outline, single_feedback)
                except ValueError as e:
                    st.error(str(e))
                    st.stop()

                with st.spinner("正在调用 DeepSeek，大概需要 30-60 秒..."):
                    try:
                        result = call_deepseek(api_key, system_prompt, user_prompt, model, temperature)
                        word_count = count_chinese_words(result)

                        st.header("3. 生成结果")
                        st.text_area("正文", result, height=500)
                        st.info(f"中文字数约：{word_count}")

                        st.download_button(
                            label="📥 下载为 .md",
                            data=result.encode("utf-8"),
                            file_name=f"不周_Ch{single_chapter}_{single_mode}.md",
                            mime="text/markdown",
                        )
                    except Exception as e:
                        st.error(f"生成失败：{e}")

        st.divider()

        # 批量改写
        st.header("3. 批量改写 / 导出")
        col_b1, col_b2, col_b3 = st.columns(3)
        with col_b1:
            batch_start = st.number_input("起始章节", min_value=1, max_value=200, value=7, step=1, key="batch_start")
        with col_b2:
            batch_end = st.number_input("结束章节", min_value=1, max_value=200, value=17, step=1, key="batch_end")
        with col_b3:
            batch_size = st.number_input("每批数量", min_value=1, max_value=5, value=3, step=1,
                                         help="Streamlit Cloud 单次运行太长会超时，建议每批 3 章")

        batch_mode = st.selectbox("批量模式", ["改写", "扩写"], index=0, key="batch_mode")

        if st.button("🚀 批量改写", type="primary", key="batch_run"):
            if not api_key:
                st.error("请先填写 DeepSeek API Key")
            elif batch_end < batch_start:
                st.error("结束章节不能小于起始章节")
            else:
                results = {}
                progress_bar = st.progress(0)
                status = st.empty()

                chapters_to_run = list(range(batch_start, batch_end + 1))
                total = len(chapters_to_run)

                for idx, ch in enumerate(chapters_to_run):
                    status.text(f"正在处理第 {ch} 章... ({idx + 1}/{total})")
                    body, outline = extract_chapter(paragraphs, ch)

                    try:
                        user_prompt = build_user_prompt(ch, batch_mode, body, outline)
                    except ValueError as e:
                        results[ch] = f"【跳过】{e}"
                        st.error(f"第 {ch} 章：{e}")
                        progress_bar.progress((idx + 1) / total)
                        continue

                    try:
                        result = call_deepseek(api_key, system_prompt, user_prompt, model, temperature)
                        results[ch] = result
                    except Exception as e:
                        results[ch] = f"【生成失败】{e}"
                        st.error(f"第 {ch} 章失败：{e}")

                    progress_bar.progress((idx + 1) / total)

                st.session_state.batch_results = results
                st.success(f"批量改写完成，共 {len(results)} 章")

        # 显示批量结果和下载
        if "batch_results" in st.session_state and st.session_state.batch_results:
            results = st.session_state.batch_results

            with st.expander("查看批量结果"):
                for ch in sorted(results.keys()):
                    st.subheader(f"第 {ch} 章")
                    st.text_area(f"第{ch}章结果", results[ch], height=300, key=f"result_ch_{ch}")

            combined_md = "\n\n---\n\n".join(
                [f"{results[ch]}" for ch in sorted(results.keys())]
            )

            col_d1, col_d2 = st.columns(2)
            with col_d1:
                st.download_button(
                    label="📥 下载全部结果.md",
                    data=combined_md.encode("utf-8"),
                    file_name=f"不周_Ch{batch_start}-{batch_end}_{batch_mode}.md",
                    mime="text/markdown",
                )
            with col_d2:
                doc = Document()
                for ch in sorted(results.keys()):
                    doc.add_paragraph(results[ch])
                    doc.add_paragraph()
                docx_io = BytesIO()
                doc.save(docx_io)
                docx_io.seek(0)
                st.download_button(
                    label="📥 下载全部结果.docx",
                    data=docx_io.getvalue(),
                    file_name=f"不周_Ch{batch_start}-{batch_end}_{batch_mode}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )

    else:
        st.info("请选择或上传 docx 文件后开始")


# ============== Tab 2: 创作大纲 ==============
with tab2:
    st.header("创作大纲")
    outline_files = [f for f in list_repo_docs() if "大纲" in f or "outline" in f.lower()]
    if outline_files:
        selected_outline = st.selectbox("选择大纲文件", outline_files, key="t2_outline")
        if selected_outline:
            content = load_repo_doc_text(selected_outline)
            st.text_area("大纲内容", content, height=600, key="t2_content")
            st.download_button(
                label="📥 下载大纲",
                data=content.encode("utf-8"),
                file_name=selected_outline,
                mime="text/plain",
            )
    else:
        st.warning("仓库 `docs/` 目录下没有找到文件名包含「大纲」的文档。请把大纲文件放进 docs/ 目录。")


# ============== Tab 3: 生成细纲 ==============
with tab3:
    st.header("根据大纲生成细纲")
    outline_files = [f for f in list_repo_docs() if "大纲" in f or "outline" in f.lower()]
    if outline_files:
        selected_outline = st.selectbox("选择大纲文件", outline_files, key="t3_outline")
        col_t3_1, col_t3_2 = st.columns(2)
        with col_t3_1:
            volume = st.number_input("卷数", min_value=1, max_value=10, value=1, step=1, key="t3_volume")
        with col_t3_2:
            chapter_count = st.number_input("本章细纲共多少章", min_value=1, max_value=200, value=60, step=1, key="t3_ch_count")

        if st.button("🚀 生成细纲", type="primary", key="t3_run"):
            if not api_key:
                st.error("请先填写 DeepSeek API Key")
            elif selected_outline:
                with st.spinner("正在读取大纲并调用 DeepSeek..."):
                    outline_content = load_repo_doc_text(selected_outline)
                    prompt = build_outline_prompt(outline_content, volume, chapter_count)
                    try:
                        result = call_deepseek(api_key, system_prompt, prompt, model, temperature)
                        st.text_area("生成结果", result, height=500, key="t3_result")
                        st.download_button(
                            label="📥 下载细纲.md",
                            data=result.encode("utf-8"),
                            file_name=f"不周_卷{volume}细纲.md",
                            mime="text/markdown",
                        )
                    except Exception as e:
                        st.error(f"生成失败：{e}")
    else:
        st.warning("仓库 `docs/` 目录下没有找到大纲文件。请把大纲文件放进 docs/ 目录。")


# ============== Tab 4: 续写 ==============
with tab4:
    st.header("1. 选择正文源文件")
    source_option_t4 = st.radio("文件来源", ["从仓库选择", "本地上传"], index=0, horizontal=True, key="t4_source")

    paragraphs_t4 = None
    if source_option_t4 == "从仓库选择":
        body_files = [f for f in list_repo_docs() if "正文" in f]
        if body_files:
            selected_body = st.selectbox("仓库正文文件", body_files, key="t4_body")
            if selected_body:
                file_bytes = load_repo_doc(selected_body)
                if file_bytes:
                    paragraphs_t4 = load_docx_text(file_bytes)
                    st.success(f"已加载正文：{selected_body}，共 {len(paragraphs_t4)} 段")
        else:
            st.warning("仓库 `docs/` 目录下没有找到文件名包含「正文」的 docx。")
    else:
        uploaded_file_t4 = st.file_uploader("上传《不周》正文 docx 文件", type=["docx"], key="t4_upload")
        if uploaded_file_t4 is not None:
            paragraphs_t4 = load_docx_text(uploaded_file_t4.getvalue())
            st.success(f"文件已读取，共 {len(paragraphs_t4)} 段")

    if paragraphs_t4 is not None:
        st.header("2. 续写设置")
        col_t4_1, col_t4_2 = st.columns(2)
        with col_t4_1:
            next_chapter = st.number_input("要续写的章节号", min_value=1, max_value=200, value=1, step=1, key="t4_ch")
        with col_t4_2:
            prev_count = st.number_input("参考前几章正文", min_value=1, max_value=10, value=3, step=1, key="t4_prev")

        use_external_outline = st.checkbox("使用独立细纲文件（如果源文件里没有细纲）", value=False, key="t4_use_outline")
        external_outline = ""
        if use_external_outline:
            outline_files_t4 = [f for f in list_repo_docs() if "细纲" in f]
            if outline_files_t4:
                selected_outline_t4 = st.selectbox("选择细纲文件", outline_files_t4, key="t4_outline")
                if selected_outline_t4:
                    external_outline = load_repo_doc_text(selected_outline_t4)
                    st.success(f"已加载细纲：{selected_outline_t4}")
            else:
                st.warning("仓库 `docs/` 目录下没有找到文件名包含「细纲」的文档。")

        if st.button("🚀 开始续写", type="primary", key="t4_run"):
            if not api_key:
                st.error("请先填写 DeepSeek API Key")
            else:
                with st.spinner("正在提取前文和细纲..."):
                    # 提取前文
                    prev_texts = []
                    start_ch = max(1, next_chapter - prev_count)
                    for ch in range(start_ch, next_chapter):
                        body, _ = extract_chapter(paragraphs_t4, ch)
                        if body.strip():
                            prev_texts.append(f"第{ch}章：\n{body}")

                    # 提取细纲
                    if use_external_outline and external_outline:
                        outline = extract_outline_from_text(external_outline, next_chapter)
                    else:
                        _, outline = extract_chapter(paragraphs_t4, next_chapter)

                    if not outline.strip():
                        st.warning("未找到本章细纲，将自由续写")

                    ch1_case = load_repo_knowledge_text("不周_Ch1_去AI味案例.md")
                    prompt = build_continuation_prompt(next_chapter, prev_texts, outline, ch1_case)

                with st.spinner("正在调用 DeepSeek 续写..."):
                    try:
                        result = call_deepseek(api_key, system_prompt, prompt, model, temperature)
                        word_count = count_chinese_words(result)

                        st.header("3. 续写结果")
                        st.text_area("正文", result, height=500, key="t4_result")
                        st.info(f"中文字数约：{word_count}")

                        st.download_button(
                            label="📥 下载为 .md",
                            data=result.encode("utf-8"),
                            file_name=f"不周_Ch{next_chapter}_续写.md",
                            mime="text/markdown",
                        )
                    except Exception as e:
                        st.error(f"续写失败：{e}")

    else:
        st.info("请选择或上传正文 docx 文件后开始")
